import docx

doc = docx.Document(r"C:\Users\lenovo\Downloads\Plantilla Tipo Tarea_CIENCIAS SOCIALES Y AMBIENTE -ACTIVIDAD-1.docx")
cell = doc.tables[1].rows[2].cells[0]

last_elm = cell._tc[-1]
print("Ultimo elemento de cell._tc:", last_elm.tag.split("}")[-1])

# If it ends with tbl, append an empty paragraph
if last_elm.tag.endswith("tbl"):
    print("Violación de esquema detectada: la celda termina en <w:tbl>. Agregando <w:p> final...")
    p = cell.add_paragraph()
    p.paragraph_format.space_before = docx.shared.Pt(0)
    p.paragraph_format.space_after = docx.shared.Pt(0)

print("Nuevo ultimo elemento:", cell._tc[-1].tag.split("}")[-1])
doc.save(r"C:\Users\lenovo\Downloads\test_repaired.docx")
