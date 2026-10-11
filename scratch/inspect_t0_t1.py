# -*- coding: utf-8 -*-
import docx

doc = docx.Document(r"C:\Users\lenovo\Downloads\Plantilla Tipo Tarea_CIENCIAS SOCIALES Y AMBIENTE -ACTIVIDAD-1.docx")

print("Header text:")
for s in doc.sections:
    for p in s.header.paragraphs:
        print(f"  '{p.text}'")

print("\nTable 0 details:")
for r_i, r in enumerate(doc.tables[0].rows):
    print(f"  Row {r_i}: {[c.text.strip().replace(chr(10), ' ') for c in r.cells]}")

print("\nTable 1 details:")
for r_i, r in enumerate(doc.tables[1].rows[:2]):
    print(f"  Row {r_i}: {[c.text.strip().replace(chr(10), ' ') for c in r.cells]}")
