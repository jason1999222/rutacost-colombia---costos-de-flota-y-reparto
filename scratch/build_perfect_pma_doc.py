# -*- coding: utf-8 -*-
"""
Generador del Documento Oficial: PLAN DE MANEJO DE IMPACTOS AMBIENTALES
Institución Universitaria Digital de Antioquia - IU Digital
Estudiantes: CARLOS ANDRÉS GARCÍA VILLADIEGO y ANDRÉS JAIR RESTREPO
Docente: ANA ELIZABETH CUARAN
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="7F7F7F", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))

def build_perfect_document():
    backup_path = r"C:\Users\lenovo\Desktop\trabajos andres\PLAN DE MANEJO DE IMPACTOS AMBIENTALES_backup.docx"
    output_docx = r"C:\Users\lenovo\Desktop\trabajos andres\PLAN DE MANEJO DE IMPACTOS AMBIENTALES.docx"
    output_pdf = r"C:\Users\lenovo\Desktop\trabajos andres\PLAN DE MANEJO DE IMPACTOS AMBIENTALES.pdf"

    print("Cargando plantilla base para preservar encabezado institucional con logo...")
    doc = docx.Document(backup_path)

    # Limpiar cuerpo manteniendo encabezado con el logo institucional
    for p in list(doc.paragraphs):
        p._p.getparent().remove(p._p)
    for t in list(doc.tables):
        t._tbl.getparent().remove(t._tbl)

    # Configuración de márgenes estándar (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Colores corporativos
    COLOR_PRIMARY = RGBColor(16, 44, 87)       # Azul oscuro institucional
    COLOR_SECONDARY = RGBColor(176, 21, 57)    # Rojo/vino tinto IU Digital
    COLOR_BODY = RGBColor(33, 33, 33)          # Gris muy oscuro legible
    HEX_HEADER_BG = "F0F4F8"                   # Fondo azul claro muy tenue para encabezados de tabla
    HEX_TITLE_BG = "E6EEF8"                    # Fondo azul suave para título de ficha

    # ========================================================
    # 1. PORTADA OFICIAL
    # ========================================================
    for _ in range(3):
        doc.add_paragraph()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("PLAN DE MANEJO DE IMPACTOS AMBIENTALES")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Formulación de Fichas de Mitigación Ambiental – Municipio de Tarazá, Antioquia\nActividad 3")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 100, 100)

    for _ in range(4):
        doc.add_paragraph()

    p_pres = doc.add_paragraph()
    p_pres.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pres = p_pres.add_run("PRESENTADO POR (TRABAJO COLABORATIVO):")
    r_pres.bold = True
    r_pres.font.name = "Arial"
    r_pres.font.size = Pt(11)
    r_pres.font.color.rgb = COLOR_SECONDARY

    p_names = doc.add_paragraph()
    p_names.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_n1 = p_names.add_run("CARLOS ANDRÉS GARCÍA VILLADIEGO\n")
    r_n1.bold = True
    r_n1.font.name = "Arial"
    r_n1.font.size = Pt(14)
    r_n1.font.color.rgb = COLOR_BODY

    r_n2 = p_names.add_run("ANDRÉS JAIR RESTREPO")
    r_n2.bold = True
    r_n2.font.name = "Arial"
    r_n2.font.size = Pt(14)
    r_n2.font.color.rgb = COLOR_BODY

    for _ in range(4):
        doc.add_paragraph()

    p_cur = doc.add_paragraph()
    p_cur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cur1 = p_cur.add_run("CURSO:\nADMINISTRACIÓN DE EMPRESAS\n\nCÁTEDRA IU DIGITAL ECOLOGÍA Y MEDIO AMBIENTE\nCÓDIGO DE GRUPO: PRECBH2602B010234\n\n")
    r_cur1.font.name = "Arial"
    r_cur1.font.size = Pt(11)
    r_cur1.bold = True

    r_doc = p_cur.add_run("DOCENTE:\nANA ELIZABETH CUARAN\n\n")
    r_doc.font.name = "Arial"
    r_doc.font.size = Pt(11)

    r_inst = p_cur.add_run("INSTITUCIÓN UNIVERSITARIA DIGITAL DE ANTIOQUIA – IU DIGITAL\nSEPTIEMBRE DE 2026")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(11)
    r_inst.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    doc.add_page_break()

    # ========================================================
    # 2. MARCO METODOLÓGICO Y ELEMENTOS DE LA FICHA
    # ========================================================
    p_h1 = doc.add_paragraph()
    r = p_h1.add_run("PLAN DE MANEJO DE IMPACTOS AMBIENTALES")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_PRIMARY

    p_met = doc.add_paragraph()
    p_met.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_met.add_run(
        "Cada uno de los impactos identificados y valorados en los ejercicios de las actividades 1 y 2 cuenta "
        "con un conjunto estructurado de acciones orientadas a la prevención, mitigación, corrección o compensación "
        "de los efectos ambientales adversos. La ficha técnica estandarizada contiene los elementos clave que "
        "facilitan la interpretación del entorno, las necesidades locales y las políticas públicas del territorio, "
        "garantizando una implementación acertada y el cumplimiento de metas acordes a la caracterización previa "
        "del municipio de Tarazá, Antioquia. Dichos elementos se definen a continuación:"
    )
    r.font.name = "Arial"
    r.font.size = Pt(10.5)

    def_items = [
        ("NOMBRE DE LA MEDIDA:", "Definición temática puntual de la acción específica propuesta."),
        ("DESCRIPCIÓN:", "Enmarca el contexto temático, conceptual, teórico y referencial que respalda la formulación de la medida."),
        ("OBJETIVO:", "Fin primordial y medible que dirige y orienta las acciones del programa."),
        ("ACTIVIDADES:", "Conjunto de tareas secuenciales y ordenadas lógicamente para la ejecución operativa."),
        ("BENEFICIOS:", "Impactos directos favorables esperados con la implementación de la medida en el medio físico, biótico o socioeconómico."),
        ("CO-BENEFICIOS:", "Impactos indirectos positivos y sinergias favorables asociadas a la medida."),
        ("ÁMBITO GEOGRÁFICO DE IMPLEMENTACIÓN:", "Localización geográfica delimitada: departamento, municipio, sector urbano, barrios y áreas de influencia directa."),
        ("ACTORES ESTRATÉGICOS:", "Entes públicos, autoridades ambientales (CORANTIOQUIA), comunidad, sector productivo y academia con corresponsabilidad en el financiamiento, ejecución y monitoreo."),
        ("ELEMENTOS HABILITANTES:", "Políticas públicas, leyes, decretos, resoluciones sectoriales y directrices del Plan de Ordenamiento Territorial (POT)."),
        ("BARRERAS PARA LA IMPLEMENTACIÓN:", "Factores antrópicos, económicos, culturales o naturales que pueden obstaculizar el cumplimiento de los objetivos."),
        ("META:", "Resultado cuantitativo y temporalmente delimitado que compromete el éxito de la intervención."),
        ("INDICADORES DE SEGUIMIENTO:", "Unidades de medición y fórmulas estandarizadas para evaluar el rendimiento, efectividad y avance de las acciones programadas.")
    ]

    for title, desc in def_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Inches(0.25)
        p_item.paragraph_format.space_after = Pt(4)
        r_t = p_item.add_run(title + " ")
        r_t.bold = True
        r_t.font.name = "Arial"
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = COLOR_PRIMARY
        r_d = p_item.add_run(desc)
        r_d.font.name = "Arial"
        r_d.font.size = Pt(10)

    doc.add_page_break()

    # ========================================================
    # 3. INTRODUCCIÓN
    # ========================================================
    p_int_title = doc.add_paragraph()
    r = p_int_title.add_run("Introducción")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_PRIMARY

    int_texts = [
        "En el marco de la gestión ambiental sostenible y el ordenamiento territorial en Colombia, la evaluación "
        "integral de las interacciones entre las actividades humanas y el medio biofísico constituye una herramienta "
        "técnica indispensable para prevenir, mitigar y controlar la degradación de los ecosistemas. En el municipio "
        "de Tarazá, situado en la subregión del Bajo Cauca antioqueño, convergen presiones antrópicas significativas "
        "derivadas del intenso tráfico vehicular sobre la Troncal de Occidente (corredor vial nacional Medellín-Costa Caribe), "
        "la expansión urbana sobre fajas de retiro, el consumo intensivo de recursos naturales no renovables y el "
        "déficit en el saneamiento hídrico básico. Esta dinámica somete a una tensión continua a los componentes atmosférico, "
        "hídrico, edáfico y biótico del territorio.",

        "El presente Plan de Manejo de Impactos Ambientales (Actividad 3) tiene como propósito formular y fundamentar "
        "técnicamente un conjunto de fichas de manejo ambiental diseñadas a partir de los hallazgos de la caracterización "
        "ecológica (Actividad 1) y la matriz de valoración de impactos (Actividad 2). Entre las alteraciones evaluadas, "
        "se priorizan aquellas calificadas con nivel de significancia «Muy Significativo (D)», tales como la contaminación "
        "del aire por material particulado (PM₁₀ y PM₂.₅), la emisión de gases de efecto invernadero (CO₂), la contaminación "
        "acústica continua por automotores pesados, la presión sobre recursos forestales y pétreos, y el vertimiento "
        "de efluentes domésticos sin tratamiento adecuado hacia las fuentes hídricas locales, con el consecuente riesgo "
        "de eutrofización y alteración de los ciclos biogeoquímicos del carbono, nitrógeno y fósforo.",

        "Para brindar una respuesta operativa, estructurada y ejecutable, cada ficha técnica delimita su marco conceptual, "
        "objetivo estratégico, ruta secuencial de actividades, beneficios directos e indirectos, delimitación geográfica barrial, "
        "corresponsabilidad de actores institucionales (liderados por CORANTIOQUIA y la Administración Municipal), marco "
        "normativo habilitante vigente en Colombia, barreras de implementación, metas cuantificables e indicadores de monitoreo "
        "con periodicidad definida. Con ello, se consolida una propuesta técnica rigurosa que armoniza el desarrollo económico "
        "con la preservación del patrimonio natural y el bienestar de la comunidad de Tarazá."
    ]

    for t in int_texts:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(t)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)

    doc.add_page_break()

    # ========================================================
    # HELPER PARA CREAR FICHAS TÉCNICAS PROFESIONALES
    # ========================================================
    def build_ficha_table(ficha_num, nombre, descripcion, objetivo, actividades,
                          beneficios, cobeneficios, ambito, actores,
                          habilitantes, barreras, metas, indicadores):

        p_fnum = doc.add_paragraph()
        r = p_fnum.add_run(f"FICHA TÉCNICA #{ficha_num}")
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(13)
        r.font.color.rgb = COLOR_SECONDARY

        table = doc.add_table(rows=8, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="8A9BA8", sz="6")

        for row in table.rows:
            make_row_cant_split(row)

        # Configurar anchos de columna: Col 0 = 3.25 in, Col 1 = 3.25 in (Total 6.5 in)
        col_widths = [Inches(3.25), Inches(3.25)]
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = col_widths[i]
                set_cell_margins(cell, top=120, bottom=120, left=140, right=140)

        # Fila 0: NOMBRE DE LA MEDIDA (Fusionada)
        cell_0 = table.cell(0, 0)
        cell_0.merge(table.cell(0, 1))
        set_cell_background(cell_0, HEX_TITLE_BG)
        p0 = cell_0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r_lbl = p0.add_run("NOMBRE DE LA MEDIDA: ")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)
        r_lbl.font.color.rgb = COLOR_PRIMARY
        r_val = p0.add_run(nombre)
        r_val.bold = True
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10.5)

        # Fila 1: DESCRIPCIÓN (Fusionada)
        cell_1 = table.cell(1, 0)
        cell_1.merge(table.cell(1, 1))
        p1 = cell_1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_lbl = p1.add_run("DESCRIPCIÓN:\n")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = COLOR_PRIMARY
        r_val = p1.add_run(descripcion)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(9.5)

        # Fila 2: OBJETIVO (Fusionada)
        cell_2 = table.cell(2, 0)
        cell_2.merge(table.cell(2, 1))
        p2 = cell_2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_lbl = p2.add_run("OBJETIVO: ")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = COLOR_PRIMARY
        r_val = p2.add_run(objetivo)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(9.5)

        # Fila 3: ACTIVIDADES (Fusionada)
        cell_3 = table.cell(3, 0)
        cell_3.merge(table.cell(3, 1))
        p3 = cell_3.paragraphs[0]
        r_lbl = p3.add_run("ACTIVIDADES SECUENCIALES:\n")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = COLOR_PRIMARY
        for idx, act in enumerate(actividades, 1):
            p_act = cell_3.add_paragraph()
            p_act.paragraph_format.left_indent = Inches(0.2)
            p_act.paragraph_format.space_after = Pt(3)
            p_act.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r_num = p_act.add_run(f"{idx}. ")
            r_num.bold = True
            r_num.font.name = "Arial"
            r_num.font.size = Pt(9.5)
            r_txt = p_act.add_run(act)
            r_txt.font.name = "Arial"
            r_txt.font.size = Pt(9.5)

        # Helper para celda de 2 columnas
        def fill_cell(cell, label, content_list, is_bullet=True):
            set_cell_background(cell, HEX_HEADER_BG)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r_l = p.add_run(label + "\n")
            r_l.bold = True
            r_l.font.name = "Arial"
            r_l.font.size = Pt(9.5)
            r_l.font.color.rgb = COLOR_PRIMARY
            set_cell_background(cell, "FFFFFF")
            if isinstance(content_list, str):
                p_c = cell.add_paragraph()
                p_c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p_c.paragraph_format.space_after = Pt(2)
                r_c = p_c.add_run(content_list)
                r_c.font.name = "Arial"
                r_c.font.size = Pt(9)
            else:
                for item in content_list:
                    p_c = cell.add_paragraph()
                    p_c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p_c.paragraph_format.space_after = Pt(2)
                    p_c.paragraph_format.left_indent = Inches(0.15)
                    bullet = "• " if is_bullet else ""
                    r_c = p_c.add_run(bullet + item)
                    r_c.font.name = "Arial"
                    r_c.font.size = Pt(9)

        # Fila 4: BENEFICIOS (Col 0) | ÁMBITO GEOGRÁFICO (Col 1)
        fill_cell(table.cell(4, 0), "BENEFICIOS DIRECTOS:", beneficios)
        fill_cell(table.cell(4, 1), "ÁMBITO GEOGRÁFICO DE IMPLEMENTACIÓN:", ambito, is_bullet=False)

        # Fila 5: CO-BENEFICIOS (Col 0) | ACTORES ESTRATÉGICOS (Col 1)
        fill_cell(table.cell(5, 0), "CO-BENEFICIOS INDIRECTOS:", cobeneficios)
        fill_cell(table.cell(5, 1), "ACTORES ESTRATÉGICOS CORRESPONSABLES:", actores)

        # Fila 6: ELEMENTOS HABILITANTES (Col 0) | BARRERAS (Col 1)
        fill_cell(table.cell(6, 0), "ELEMENTOS HABILITANTES (NORMATIVA Y POLÍTICAS):", habilitantes)
        fill_cell(table.cell(6, 1), "BARRERAS PARA LA IMPLEMENTACIÓN:", barreras)

        # Fila 7: METAS (Col 0) | INDICADORES (Col 1)
        fill_cell(table.cell(7, 0), "METAS CUANTIFICABLES:", metas)
        fill_cell(table.cell(7, 1), "INDICADORES DE MONITOREO Y SEGUIMIENTO:", indicadores)

        doc.add_page_break()

    # ========================================================
    # 4. FICHA TÉCNICA #1: AIRE Y RUIDO VIAL (TRONCAL MEDELLÍN-COSTA)
    # ========================================================
    build_ficha_table(
        ficha_num="1",
        nombre="Programa de Mitigación de Emisiones Atmosféricas y Amortiguamiento Acústico en el Corredor Vial Urbano Medellín-Costa.",
        descripcion=(
            "Esta medida responde a los impactos ambientales evaluados con calificación «Muy Significativo (D)», "
            "originados por la circulación continua de vehículos de carga pesada, transporte intermunicipal y automotores "
            "particulares sobre la Troncal de Occidente a su paso por el perímetro urbano de Tarazá. La propuesta se fundamenta "
            "en la ingeniería ecológica y las Soluciones Basadas en la Naturaleza (SbN), utilizando barreras vegetales multiestrato "
            "como filtros biológicos capaces de interceptar material particulado (PM₁₀ y PM₂.₅), adsorber gases de combustión (CO, NOₓ), "
            "fijar dióxido de carbono (CO₂) y disipar la energía de las ondas sonoras vehiculares mediante mecanismos de dispersión, "
            "difracción y absorción acústica foliar."
        ),
        objetivo=(
            "Prevenir y mitigar los impactos sobre la calidad del aire y los niveles de ruido ambiental en las zonas residenciales "
            "y comerciales adyacentes a la vía Medellín-Costa en Tarazá, mediante el establecimiento de barreras vegetales protectoras "
            "y la optimización del arbolado urbano protector."
        ),
        actividades=[
            "Realizar un levantamiento topográfico y diagnóstico microclimático para delimitar los tramos críticos del corredor vial con mayor exposición comunitaria a ruido y material particulado.",
            "Diseñar técnicamente las franjas de amortiguamiento e implantar barreras vegetales multiestrato, combinando especies arbustivas densas (ej. San Joaquín, Swinglea) con árboles nativos de follaje perenne tolerantes a emisiones vehiculares.",
            "Implementar un plan de manejo silvicultural, riego tecnificado no potable y control fitosanitario durante los primeros 24 meses para asegurar el prendimiento y vigor vegetativo.",
            "Desarrollar jornadas de concertación y cultura ambiental con comerciantes frentistas y juntas de acción comunal para la corresponsabilidad y cuidado de las coberturas verdes protectoras.",
            "Ejecutar campañas de monitoreo semestral de calidad del aire y sonometría ambiental diurna/nocturna conforme a protocolos oficiales."
        ],
        beneficios=[
            "Retención mecánica y absorción foliar de material particulado (PM₁₀ y PM₂.₅) y gases contaminantes a nivel de la faja de respiración comunitaria.",
            "Atenuación comprobada del nivel de presión sonora ambiental entre 4 dB(A) y 8 dB(A) en las fachadas de viviendas directamente adyacentes a la calzada.",
            "Captura y secuestro biológico de carbono (CO₂), aportando a la regulación microclimática y mitigación del cambio climático local."
        ],
        ambito=(
            "Departamento de Antioquia, Municipio de Tarazá, cabecera urbana. Corredor longitudinal de la Troncal de Occidente "
            "(Ruta Nacional 25 Medellín-Costa Caribe), tramo urbano de 3.2 kilómetros, con intervención prioritaria en los barrios "
            "San Jerónimo, San Nicolás, El Triunfo, Piedras y el sector comercial de La Bomba."
        ),
        cobeneficios=[
            "Reducción en la tasa de incidencia de infecciones respiratorias agudas (IRA) y fatiga auditiva en la población vulnerable infantil y de la tercera edad.",
            "Mitigación del efecto de 'isla de calor urbana' gracias a la evapotranspiración y sombra proyectada por las coberturas vegetales.",
            "Embellecimiento paisajístico y valorización socioeconómica de los predios frentistas al corredor vial."
        ],
        actores=[
            "Corporación Autónoma Regional del Centro de Antioquia (CORANTIOQUIA - Territorial Panzenú): Asesoría técnica, licenciamiento y seguimiento a la calidad del aire.",
            "Alcaldía Municipal de Tarazá (Secretaría de Planeación y Obras Públicas): Autorización de intervención de espacio público y mantenimiento agronómico.",
            "Instituto Nacional de Vías (INVIAS) / Concesión Vial encargada: Concertación del uso de fajas de retiro y señalización vial preventiva.",
            "Institución Universitaria Digital de Antioquia (IU Digital): Soporte de investigación aplicada, modelación ambiental y monitoreo participativo.",
            "Juntas de Acción Comunal (JAC) y comerciantes del corredor vial: Custodia ciudadana de las especies sembradas."
        ],
        habilitantes=[
            "Ley 99 de 1993 (Fundamentos de la política ambiental colombiana).",
            "Resolución 2254 de 2017 del MADS (Norma nacional de calidad del aire y límites permisibles).",
            "Resolución 627 de 2006 del MAVDT (Norma nacional de emisión de ruido y ruido ambiental).",
            "Ley 2169 de 2021 (Ley de Acción Climática de Colombia hacia la carbono-neutralidad).",
            "Plan de Ordenamiento Territorial (POT) y Plan de Desarrollo Municipal de Tarazá."
        ],
        barreras=[
            "Flujo vehicular pesado ininterrumpido que genera estrés mecánico y depósito continuo de hollín sobre plántulas en etapa inicial.",
            "Restricción de espacio público en bermas por invasión u ocupación comercial informal.",
            "Riesgo de déficit hídrico prolongado durante la temporada seca de la región del Bajo Cauca.",
            "Falta de apropiación inicial por parte de algunos frentistas que priorizan visibilidad comercial."
        ],
        metas=[
            "Establecer y consolidar 3.0 kilómetros lineales de barreras vivas multiestrato en el área urbana de la vía nacional en un plazo de 18 meses.",
            "Lograr una reducción de al menos 5 dB(A) en los niveles de ruido diurno registrados en receptores críticos habitacionales al finalizar el segundo año de implementación.",
            "Garantizar una tasa de sobrevivencia vegetal superior al 85% al término de los 24 meses de mantenimiento silvicultural continuo."
        ],
        indicadores=[
            "Tasa de Sobrevivencia Vegetal (%) = (Individuos vegetales vivos / Total sembrados) × 100. Meta: ≥ 85%.",
            "Nivel de Presión Sonora Continuo Equivalente diurno (LAeq en dB(A)), medido semestralmente conforme a la Res. 627/2006.",
            "Concentración de Material Particulado PM₁₀ y PM₂.₅ (µg/m³), contrastando línea base vs. zonas con barrera consolidada.",
            "Número de metros lineales de barrera viva implantada y mantenida con acta de entrega técnica."
        ]
    )

    # ========================================================
    # 5. FICHA TÉCNICA #2: MATERIALES Y ECONOMÍA CIRCULAR (RCD Y MADERA)
    # ========================================================
    build_ficha_table(
        ficha_num="2",
        nombre="Programa de Optimización de Materiales, Economía Circular y Gestión Integral de Residuos de Construcción y Demolición (RCD).",
        descripcion=(
            "Esta medida atiende la problemática de agotamiento de materias primas no renovables (agregados pétreos, arenas de río "
            "y arcillas) y la deforestación ocasionada por el consumo de madera sin certificación legal en obras civiles y carpintería "
            "doméstica en Tarazá. La propuesta se fundamenta en el marco teórico de la Economía Circular y el Análisis de Ciclo de Vida (ACV), "
            "fomentando la prolongación de la vida útil de los materiales, la sustitución ecoeficiente, la valorización de escombros de obra "
            "y la adquisición responsable conforme a la política nacional de compras públicas sostenibles."
        ),
        objetivo=(
            "Disminuir la presión extractiva sobre recursos naturales maderables y pétreos en el municipio de Tarazá, promoviendo "
            "el reaprovechamiento de materiales de construcción, la reutilización de elementos de madera y la certificación de origen legal."
        ),
        actividades=[
            "Elaborar y socializar una guía técnica municipal de eco-materiales, compras sostenibles y criterios de verificación de salvoconductos de movilización de madera legal.",
            "Desarrollar talleres comunitarios de capacitación en técnicas de carpintería circular, restauración de muebles y preservación no tóxica de maderas para alargar su ciclo útil.",
            "Concertar con depósitos y ferreterías de Tarazá la habilitación de líneas de comercialización de insumos certificados y agregados reciclados.",
            "Crear un Punto Limpio de Acopio Temporal y Transferencia para Residuos de Construcción y Demolición (RCD), clasificando sobrantes limpios (ladrillo, concreto, mampostería) para sub-bases de vías terciarias y obras comunitarias.",
            "Establecer incentivos no tributarios y distinciones de 'Negocio Verde Local' a comercios comprometidos con la economía circular."
        ],
        beneficios=[
            "Disminución directa en la tasa de extracción de materiales de canteras y lechos de ríos en la cuenca del Río Tarazá.",
            "Reducción de la deforestación local al desincentivar la compra de madera nativa de procedencia ilícita.",
            "Desvío de escombros de botaderos clandestinos, orillas de fuentes hídricas y espacio público municipal."
        ],
        ambito=(
            "Departamento de Antioquia, Municipio de Tarazá, casco urbano y periferia periurbana, focalizado en ferreterías, "
            "aserríos, talleres de carpintería y proyectos de construcción en los barrios Las Palmas, Villa Esperanza y San Jerónimo."
        ),
        cobeneficios=[
            "Ahorro económico directo para familias vulnerables al autogestionar el mantenimiento de estructuras y mobiliario.",
            "Disminución de costos municipales por recolección y disposición de escombros arrojados clandestinamente.",
            "Generación de encadenamientos productivos y empleo verde informal en el reciclaje y transformación de RCD."
        ],
        actores=[
            "Corporación Autónoma Regional del Centro de Antioquia (CORANTIOQUIA - Territorial Panzenú): Control al tráfico ilegal de madera, licenciamiento y aplicación de la Res. 0472/2017.",
            "Alcaldía Municipal de Tarazá (Secretaría de Planeación e Infraestructura): Destinación del lote para el Centro de Acopio de RCD y regulación municipal.",
            "Ferreterías, depósitos de materiales y talleres de ebanistería de Tarazá: Adopción de buenas prácticas comerciales.",
            "Institución Universitaria Digital de Antioquia (IU Digital): Capacitación técnica comunitaria y diseño de indicadores circulares.",
            "Comunidad y maestros de obra locales: Separación en fuente y entrega de sobrantes en puntos autorizados."
        ],
        habilitantes=[
            "Resolución 0472 de 2017 del MADS (Reglamentación integral de Residuos de Construcción y Demolición - RCD).",
            "Estrategia Nacional de Economía Circular (ENEC) del Gobierno de Colombia.",
            "Pacto Intersectorial por la Madera Legal en Colombia (MinAmbiente y CORANTIOQUIA).",
            "Ley 99 de 1993 y Código Nacional de Recursos Naturales Renovables (Decreto Ley 2811 de 1974).",
            "Plan de Gestión Integral de Residuos Sólidos (PGIRS) del Municipio de Tarazá."
        ],
        barreras=[
            "Prevalencia cultural de técnicas constructivas tradicionales sin criterios de reciclaje o separación de escombros.",
            "Precios más bajos de la madera no certificada extraída informalmente en zonas rurales apartadas.",
            "Capacidad técnica y logística limitada del municipio para la trituración mecánica y valorización a gran escala de RCD.",
            "Falta de espacio adecuado para disposición temporal de materiales en obras de pequeña escala."
        ],
        metas=[
            "Lograr que al cabo de 24 meses el 45% de los proyectos de remodelación y construcción menor en el área urbana incorporen materiales recuperados o RCD reciclados.",
            "Reducir en un 30% el volumen de escombros de construcción arrojados en espacios públicos y rondas hídricas de Tarazá.",
            "Capacitar a un mínimo de 150 maestros de obra, carpinteros y líderes comunitarios en economía circular y aprovechamiento de materiales sostenibles."
        ],
        indicadores=[
            "Porcentaje de aprovechamiento de RCD (%) = (Toneladas de RCD valorizadas / Total de RCD generados en el municipio) × 100.",
            "Volumen de madera con salvoconducto de movilización y certificación legal comercializada en ferreterías locales (m³/año).",
            "Número de personas capacitadas y certificadas en talleres de carpintería circular y ecoconstrucción.",
            "Número de depósitos de materiales articulados formalmente al Punto Limpio de Acopio Municipal."
        ]
    )

    # ========================================================
    # 6. FICHA TÉCNICA #3: USO EFICIENTE DEL AGUA Y CONTROL DE VERTIMIENTOS
    # ========================================================
    build_ficha_table(
        ficha_num="3",
        nombre="Programa de Uso Eficiente del Agua y Control Integral de Vertimientos de Aguas Residuales Domésticas en Microcuencas Urbanas.",
        descripcion=(
            "Esta medida mitiga el agotamiento del recurso hídrico superficial por consumo doméstico excesivo y previene el deterioro "
            "ecológico provocado por las descargas de aguas residuales sin tratamiento previo en las fuentes hídricas de Tarazá. La propuesta "
            "se sustenta teóricamente en la Gestión Integral del Recurso Hídrico (GIRH) y el saneamiento ecológico descentralizado, "
            "incorporando tecnologías de bajo costo para el ahorro de agua potable, cosecha de aguas lluvias y sistemas biológicos "
            "de retención de grasas y nutrientes para salvaguardar los ciclos biogeoquímicos del nitrógeno y fósforo, evitando la eutrofización."
        ),
        objetivo=(
            "Optimizar el consumo per cápita de agua potable en los hogares urbanos de Tarazá y reducir la carga contaminante orgánica "
            "y química vertida directamente a las microcuencas receptoras de la cuenca del Río Tarazá."
        ),
        actividades=[
            "Instalar kits demostrativos de dispositivos ahorradores de agua (aireadores de grifería, reductores de caudal y sanitarios de bajo volumen) en viviendas piloto e instituciones educativas.",
            "Diseñar e instalar Sistemas de Cosecha y Aprovechamiento de Agua Lluvia (SCALL) para usos no potables (riego de jardines, aseo exterior y sanitarios).",
            "Desarrollar talleres comunitarios sobre gestión de sustancias nocivas en el hogar, prohibiendo el vertimiento de aceites de cocina usados y químicos a los desagües.",
            "Promover e implementar trampas de grasa domiciliarias y pozos sépticos con filtros anaeróbicos de flujo ascendente (FAFA) en sectores sin cobertura de alcantarillado.",
            "Efectuar monitoreos participativos periódicos de calidad físico-química (DBO₅, DQO, Grasas/Aceites, SST) en puntos de vertimiento a quebradas urbanas."
        ],
        beneficios=[
            "Reducción directa en la tasa de extracción de agua cruda en las bocatomas del acueducto municipal.",
            "Disminución drástica de la carga de materia orgánica (DBO₅), aceites y sólidos suspendidos totales vertidos a las fuentes hídricas.",
            "Prevención de fenómenos de anoxia y eutrofización en los tramos lentos de las microcuencas receptoras."
        ],
        ambito=(
            "Departamento de Antioquia, Municipio de Tarazá, sector urbano residencial, priorizando los barrios colindantes con "
            "la Quebrada La Cieza, Quebrada La Lucha y la margen receptora izquierda del Río Tarazá (barrios El Bosque y El Carmen)."
        ),
        cobeneficios=[
            "Disminución significativa en el valor de la factura de acueducto para los hogares por reducción de metros cúbicos consumidos.",
            "Control y mitigación de malos olores y erradicación de vectores biológicos de enfermedades (zancudos, roedores).",
            "Recuperación del valor paisajístico y recreativo de los cuerpos de agua urbanos para la comunidad ribereña."
        ],
        actores=[
            "Empresa de Servicios Públicos Domiciliarios de Tarazá: Soporte en medición de micromedición, suministro y gestión de alcantarillado.",
            "Corporación Autónoma Regional del Centro de Antioquia (CORANTIOQUIA - Territorial Panzenú): Control de permisos de vertimiento y seguimiento analítico de calidad de agua.",
            "Alcaldía Municipal de Tarazá (Secretaría de Salud y Bienestar Social): Vigilancia epidemiológica y educación sanitaria.",
            "Institución Universitaria Digital de Antioquia (IU Digital): Análisis de laboratorio, capacitación técnica y evaluación de eficiencia de trampas.",
            "Comunidad barrial y líderes de acueductos comunitarios: Adopción de buenas prácticas y mantenimiento de trampas de grasa."
        ],
        habilitantes=[
            "Ley 373 de 1997 (Programa para el Uso Eficiente y Ahorro del Agua - PUEAA).",
            "Decreto 1076 de 2015 (Reglamentario del Sector Ambiente y Desarrollo Sostenible).",
            "Resolución 0631 de 2015 del MADS (Parámetros y valores límites máximos permisibles en vertimientos puntuales a aguas superficiales).",
            "Resolución 1256 de 2021 del MADS (Lineamientos técnicos para el reúso de aguas residuales tratadas y aguas lluvias).",
            "Objetivo de Desarrollo Sostenible ODS 6 (Agua limpia y saneamiento universal)."
        ],
        barreras=[
            "Cobertura parcial e incompleta de la red de alcantarillado sanitario en barrios periféricos y de invasión reciente.",
            "Costos iniciales percibidos para la compra de tanques de recolección y tuberías de sistemas de aguas lluvias.",
            "Cultura arraigada de vertimiento de aceites de freír y detergentes fosfatados directamente al sifón.",
            "Alta turbiedad y sedimentación del agua en épocas de crecientes súbitas invernales que colmata sistemas básicos."
        ],
        metas=[
            "Disminuir en un 18% el consumo promedio per cápita de agua potable en los hogares intervenidos en un lapso de 24 meses.",
            "Lograr que el 75% de los hogares sin conexión al alcantarillado cuenten con trampa de grasas y biofiltro antes de descargar aguas a los afluentes.",
            "Instalar 50 sistemas de cosecha de agua de lluvia (SCALL) en edificaciones comunitarias y viviendas del casco urbano en 2 años."
        ],
        indicadores=[
            "Consumo medio mensual por suscriptor (m³/vivienda/mes), verificado a través de registros de micromedición del operador.",
            "Porcentaje de reducción de carga orgánica DBO₅ y Grasas/Aceites (%) en los puntos de vertimiento intervenidos.",
            "Número de trampas de grasa domiciliarias y biofiltros instalados, operando y verificados técnicamente.",
            "Volumen estimado de agua lluvia captada y reutilizada para fines no potables (m³/año)."
        ]
    )

    # ========================================================
    # 7. CONCLUSIONES SOBRE LOS HALLAZGOS (25 PUNTOS DE LA RÚBRICA)
    # ========================================================
    p_con_title = doc.add_paragraph()
    r = p_con_title.add_run("CONCLUSIÓN SOBRE LOS HALLAZGOS Y RELACIÓN DE MEDIDAS")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_PRIMARY

    con_texts = [
        "El desarrollo metódico de este estudio y la subsiguiente estructuración de las fichas de manejo ambiental permiten "
        "establecer un conjunto de conclusiones estratégicas sobre la interacción socioecológica en el municipio de Tarazá, "
        "demostrando una relación directa, manifiesta y causal entre los impactos identificados en la caracterización (Actividad 1), "
        "evaluados en la matriz (Actividad 2), y las medidas de manejo formuladas para su intervención efectiva:",

        "1. Correspondencia directa entre impactos críticos y medidas de mitigación:\n"
        "La matriz de valoración ambiental evidenció que la mayor presión sobre el entorno urbano de Tarazá proviene de la operación "
        "continua del corredor vial Troncal de Occidente (Medellín-Costa Caribe), calificado con un nivel de significancia «Muy Significativo (D)» "
        "debido a la emisión masiva de gases de combustión, material particulado y contaminación acústica sostenida. Frente a este hallazgo, "
        "la Ficha Técnica #1 responde con una medida fundamentada en Soluciones Basadas en la Naturaleza (SbN), utilizando barreras vivas multiestrato "
        "que actúan físicamente como pantallas deflectoras del ruido y biológicamente como trampas de retención de aerosoles y sumideros de carbono. "
        "De igual forma, la presión extractiva sobre canteras y bosques para madera y mampostería se mitiga de manera directa en la Ficha Técnica #2 "
        "mediante la economía circular y la valorización de escombros (RCD), mientras que el riesgo de eutrofización y degradación hídrica "
        "en las microcuencas receptoras es neutralizado puntualmente por la Ficha Técnica #3 mediante dispositivos de ahorro, trampas de grasas "
        "y biofiltros descentralizados.",

        "2. Aplicación rigurosa de la Jerarquía de Mitigación Ambiental:\n"
        "En estricto cumplimiento de los principios rectores de la gestión ambiental moderna y las competencias evaluadas en la asignatura, "
        "el conjunto de fichas técnicas articula de manera integral y explícita los cuatro niveles jerárquicos de intervención:\n"
        "• Medidas de PREVENCIÓN: Orientadas a evitar que el impacto adverso ocurra. Se plasman en la capacitación comunitaria para el consumo "
        "responsable de agua, las guías de compras públicas sostenibles que impiden la adquisición de madera ilegal no certificada, y la prohibición "
        "estricta de vertimiento de aceites de cocina al alcantarillado.\n"
        "• Medidas de MITIGACIÓN: Diseñadas para reducir la magnitud, duración o severidad del impacto inevitable. Se materializan en la disipación "
        "del ruido vial mediante follaje denso, la absorción foliar de material particulado y la disminución de caudales extraídos mediante griferías ecoeficientes.\n"
        "• Medidas de CORRECCIÓN: Encaminadas a tratar, recuperar o modificar el componente alterado. Se evidencian en la retención física de grasas "
        "y aceites mediante biofiltros antes de su salida a cuerpos de agua, así como en la trituración, clasificación y reciclaje de RCD en el Punto Limpio municipal.\n"
        "• Medidas de COMPENSACIÓN: Concebidas para resarcir la pérdida biótica y climática residual que no puede ser totalmente prevenida ni mitigada. "
        "Se ejecutan a través de la siembra masiva de especies florísticas y medicinales nativas en el espacio público urbano de Tarazá, expandiendo "
        "la biomasa vegetal como sumidero activo de carbono y regulador del ciclo local del oxígeno.",

        "3. Rol estratégico del componente biótico y equilibrio de ciclos biogeoquímicos:\n"
        "El análisis reafirma que el patrimonio vegetal de Tarazá no representa un elemento pasivo, sino un motor activo de resiliencia ecológica. "
        "La integración planificada de vegetación protectora a lo largo del eje vial y las cuencas urbanas salvaguarda el equilibrio de los ciclos "
        "biogeoquímicos locales: el ciclo del carbono y del oxígeno mediante la fotosíntesis y captura biológica, y los ciclos del nitrógeno y fósforo "
        "al prevenir la llegada de descargas orgánicas directas a los afluentes del Río Tarazá y Río Cauca, evitando procesos catastróficos de eutrofización.",

        "4. Gobernanza ambiental, corresponsabilidad y factibilidad territorial:\n"
        "Finalmente, la viabilidad de este Plan de Manejo radica en su articulación sinérgica entre la institucionalidad ambiental (CORANTIOQUIA - Territorial Panzenú), "
        "la gestión de la Administración Municipal de Tarazá, el sector comercial y la ciudadanía organizada en Juntas de Acción Comunal, contando con el "
        "acompañamiento técnico, investigativo y pedagógico de la Institución Universitaria Digital de Antioquia (IU Digital). Esta conjunción asegura que "
        "las metas e indicadores formulados no queden como meras declaraciones teóricas, sino que se transformen en instrumentos vinculantes de desarrollo "
        "sostenible y ordenamiento territorial para el Bajo Cauca antioqueño."
    ]

    for t in con_texts:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(t)
        r.font.name = "Arial"
        r.font.size = Pt(10)

    doc.add_page_break()

    # ========================================================
    # 8. REFERENCIAS BIBLIOGRÁFICAS Y NORMATIVAS
    # ========================================================
    p_ref_title = doc.add_paragraph()
    r = p_ref_title.add_run("REFERENCIAS BIBLIOGRÁFICAS Y MARCO NORMATIVO")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_PRIMARY

    refs = [
        "Congreso de la República de Colombia. (1993). Ley 99 de 1993: Por la cual se crea el Ministerio del Medio Ambiente, se reordena el Sector Público encargado de la gestión y conservación del medio ambiente y los recursos naturales renovables, se organiza el Sistema Nacional Ambiental, SINA y se dictan otras disposiciones. Diario Oficial No. 41.146.",
        "Congreso de la República de Colombia. (1997). Ley 373 de 1997: Por la cual se establece el programa para el uso eficiente y ahorro del agua. Diario Oficial No. 43.058.",
        "Congreso de la República de Colombia. (2021). Ley 2169 de 2021: Por medio de la cual se impulsa el desarrollo bajo en carbono del país mediante el establecimiento de metas y medidas mínimas en materia de carbono neutralidad y resiliencia climática y se dictan otras disposiciones (Ley de Acción Climática).",
        "Corporación Autónoma Regional del Centro de Antioquia - CORANTIOQUIA. (2020). Plan de Gestión Ambiental Regional (PGAR) 2020-2031: Territorios de vida, sostenibilidad y resiliencia en la jurisdicción del Centro de Antioquia. Medellín, Colombia.",
        "Ministerio de Ambiente y Desarrollo Sostenible - MADS. (2015). Decreto 1076 de 2015: Por medio del cual se expide el Decreto Único Reglamentario del Sector Ambiente y Desarrollo Sostenible. Bogotá D.C., Colombia.",
        "Ministerio de Ambiente y Desarrollo Sostenible - MADS. (2015). Resolución 0631 de 2015: Por la cual se establecen los parámetros y los valores límites máximos permisibles en los vertimientos puntuales a cuerpos de aguas superficiales y a los sistemas de alcantarillado público. Bogotá D.C.",
        "Ministerio de Ambiente y Desarrollo Sostenible - MADS. (2017). Resolución 2254 de 2017: Por la cual se adopta la norma de calidad del aire ambiente y se dictan otras disposiciones. Bogotá D.C.",
        "Ministerio de Ambiente y Desarrollo Sostenible - MADS. (2017). Resolución 0472 de 2017: Por la cual se reglamenta la gestión integral de los residuos generados en las actividades de construcción y demolición (RCD) y se dictan otras disposiciones.",
        "Ministerio de Ambiente y Desarrollo Sostenible - MADS. (2018). Estrategia Nacional de Economía Circular: Cierre de ciclos de materiales, innovación tecnológica, colaboración y nuevos modelos de negocio. Gobierno de Colombia.",
        "Ministerio de Ambiente y Desarrollo Sostenible - MADS. (2021). Resolución 1256 de 2021: Por la cual se reglamenta el uso de aguas residuales tratadas y aguas lluvias y se adoptan otras disposiciones. Bogotá D.C.",
        "Ministerio de Ambiente, Vivienda y Desarrollo Territorial - MAVDT. (2006). Resolución 627 de 2006: Por la cual se establece la norma nacional de emisión de ruido y ruido ambiental. Bogotá D.C.",
        "Municipio de Tarazá. (2020). Plan de Desarrollo Municipal de Tarazá 2020-2023 / 2024-2027: Hacia la consolidación del territorio, la paz y la sostenibilidad ambiental. Alcaldía Municipal de Tarazá, Antioquia."
    ]

    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_ref.paragraph_format.left_indent = Inches(0.3)
        p_ref.paragraph_format.first_line_indent = Inches(-0.3)
        p_ref.paragraph_format.space_after = Pt(6)
        p_ref.paragraph_format.line_spacing = 1.15
        r_ref = p_ref.add_run(ref)
        r_ref.font.name = "Arial"
        r_ref.font.size = Pt(9)

    doc.save(output_docx)
    print("Documento DOCX generado exitosamente en:", output_docx)

if __name__ == "__main__":
    build_perfect_document()
