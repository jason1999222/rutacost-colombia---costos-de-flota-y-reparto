# -*- coding: utf-8 -*-
"""
Script generador del documento oficial de la PRIMERA PARTE DEL ANTEPROYECTO DE INVESTIGACIÓN / MONOGRAFÍA.
Formato institucional: Decanatura Académica - Facultad de Psicología - Bachelor of Arts
Código oficial: Formato F_CA_0013, Versión 3, Fecha 20_09_2024
Estudiantes: Bertha Inés Medina Hernández y Yudy Carolina Benítez
Título: Impactos psicológicos del trabajo sexual en la zona de tolerancia de Medellín:
        Una revisión de las políticas públicas gubernamentales entre el periodo 2020 - 2024.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))

def add_header_table(header):
    table = header.add_table(rows=1, cols=2, width=Inches(6.8))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="000000", sz="6")
    
    cell_left = table.rows[0].cells[0]
    cell_left.width = Inches(4.7)
    set_cell_margins(cell_left, top=100, bottom=100, left=120, right=120)
    
    p0 = cell_left.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(1)
    p0.paragraph_format.line_spacing = 1.05
    r0 = p0.add_run("Decanatura Académica")
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(11)
    
    p1 = cell_left.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(1)
    p1.paragraph_format.space_after = Pt(1)
    p1.paragraph_format.line_spacing = 1.05
    r1 = p1.add_run("Facultad de psicología")
    r1.bold = True
    r1.italic = True
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    
    p2 = cell_left.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(1)
    p2.paragraph_format.space_after = Pt(1)
    p2.paragraph_format.line_spacing = 1.05
    r2 = p2.add_run("Bachelor of arts")
    r2.bold = True
    r2.italic = True
    r2.font.name = "Arial"
    r2.font.size = Pt(10.5)
    
    p3 = cell_left.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(1)
    p3.paragraph_format.space_after = Pt(2)
    p3.paragraph_format.line_spacing = 1.05
    r3 = p3.add_run("formato de trabajo de Grado/Investigación")
    r3.bold = True
    r3.italic = True
    r3.font.name = "Arial"
    r3.font.size = Pt(11)
    
    cell_right = table.rows[0].cells[1]
    cell_right.width = Inches(2.1)
    set_cell_margins(cell_right, top=40, bottom=40, left=60, right=60)
    
    subtable = cell_right.add_table(rows=5, cols=2)
    subtable.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(subtable, color="000000", sz="4")
    
    cell_logo = subtable.rows[0].cells[0]
    cell_logo.merge(subtable.rows[0].cells[1])
    set_cell_margins(cell_logo, top=40, bottom=40, left=40, right=40)
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(0)
    r_logo = p_logo.add_run("ΣIU  |  IU Digital")
    r_logo.bold = True
    r_logo.font.name = "Arial"
    r_logo.font.size = Pt(9.5)
    r_logo.font.color.rgb = RGBColor(16, 44, 87)
    
    meta_rows = [
        ("Formato", "F_CA_0013"),
        ("Versión", "3"),
        ("Fecha", "20_09_2024")
    ]
    for idx, (label, val) in enumerate(meta_rows, start=1):
        c_lbl = subtable.rows[idx].cells[0]
        c_lbl.width = Inches(0.9)
        set_cell_margins(c_lbl, top=20, bottom=20, left=40, right=40)
        p = c_lbl.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8)
        
        c_val = subtable.rows[idx].cells[1]
        c_val.width = Inches(1.1)
        set_cell_margins(c_val, top=20, bottom=20, left=40, right=40)
        pv = c_val.paragraphs[0]
        pv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pv.paragraph_format.space_before = Pt(0)
        pv.paragraph_format.space_after = Pt(0)
        rv = pv.add_run(val)
        rv.font.name = "Arial"
        rv.font.size = Pt(8)
    
    c_warn = subtable.rows[4].cells[0]
    c_warn.merge(subtable.rows[4].cells[1])
    set_cell_margins(c_warn, top=20, bottom=20, left=20, right=20)
    pw = c_warn.paragraphs[0]
    pw.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pw.paragraph_format.space_before = Pt(0)
    pw.paragraph_format.space_after = Pt(0)
    rw = pw.add_run("No Modificar esta fecha")
    rw.font.name = "Arial"
    rw.font.size = Pt(6.5)
    rw.italic = True
    
    if len(cell_right.paragraphs) > 0 and not cell_right.paragraphs[0].text.strip():
        p_elem = cell_right.paragraphs[0]._p
        p_elem.getparent().remove(p_elem)

def add_title_field(doc, label, value, is_bold_val=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r_lbl = p.add_run(label + "\n")
    r_lbl.bold = True
    r_lbl.font.name = "Arial"
    r_lbl.font.size = Pt(11)
    
    r_val = p.add_run(value)
    r_val.bold = is_bold_val
    r_val.font.name = "Arial"
    r_val.font.size = Pt(11)

def add_section_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.15
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(11.5)
        r.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 3:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.bold = True
        r.italic = True
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_p(doc, text, first_line_indent=0.5, space_after=6, italic_prefix="", bold_prefix=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(first_line_indent)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.bold = True
        rb.font.name = "Arial"
        rb.font.size = Pt(10.5)
        rb.font.color.rgb = RGBColor(0, 0, 0)
        
    if italic_prefix:
        ri = p.add_run(italic_prefix)
        ri.italic = True
        ri.font.name = "Arial"
        ri.font.size = Pt(10.5)
        ri.font.color.rgb = RGBColor(0, 0, 0)
        
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_reference_p(doc, full_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    
    r = p.add_run(full_text)
    r.font.name = "Arial"
    r.font.size = Pt(10)
    return p

def add_bullet_item(doc, text, level=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    indent_val = 0.5 + (level * 0.25)
    p.paragraph_format.left_indent = Inches(indent_val)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    r_bullet = p.add_run("●  ")
    r_bullet.font.name = "Arial"
    r_bullet.font.size = Pt(8.5)
    r_bullet.font.color.rgb = RGBColor(16, 44, 87)
    
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_index_table(doc):
    items = [
        ("1. Índice General", "1", True),
        ("2. Resumen en Español y Abstract en Inglés", "3", True),
        ("    2.1. Resumen y Palabras Clave", "3", False),
        ("    2.2. Abstract and Keywords", "3", False),
        ("3. Introducción", "5", True),
        ("4. Tema de Investigación y Delimitación Temática", "6", True),
        ("5. Planteamiento del Problema de Investigación", "7", True),
        ("    5.1. Contextualización del Trabajo Sexual en la Comuna 10 (Medellín)", "7", False),
        ("    5.2. Impactos Psicosociales y Afectaciones Clínicas en Salud Mental", "7", False),
        ("    5.3. Análisis Crítico de las Políticas Públicas (2020–2024) y Brecha Institucional", "8", False),
        ("    5.4. Formulación de la Pregunta de Investigación General y Específicas", "9", False),
        ("6. Objetivos de Investigación", "10", True),
        ("    6.1. Objetivo General", "10", False),
        ("    6.2. Objetivos Específicos", "10", False),
        ("7. Marco Teórico y Conceptual (Revisado y Consolidado)", "11", True),
        ("    7.1. Antecedentes de Investigación y Estado del Arte", "11", False),
        ("    7.2. Teoría del Estigma y Manejo de la Identidad Deteriorada (Erving Goffman)", "11", False),
        ("    7.3. Condiciones de Vida, Violencia Estructural y Precariedad Cotidiana", "12", False),
        ("    7.4. Modelo Transaccional del Estrés y Procesos de Afrontamiento (Lazarus & Folkman)", "13", False),
        ("    7.5. Resiliencia, Redes de Solidaridad entre Pares y Factores Protectores", "13", False),
        ("    7.6. Marco Contextual, Normativo y de Políticas Públicas Gubernamentales (2020–2024)", "14", False),
        ("    7.7. Sistema de Variables y Matriz de Categorías de Análisis", "14", False),
        ("9. Metodología (*)", "16", True),
        ("    9.1 Diseño de la investigación", "16", False),
        ("    9.2 Fuentes de información", "17", False),
        ("    9.3 Criterios de inclusión y exclusión", "18", False),
        ("    9.4 Técnicas de recolección y análisis de la información", "18", False),
        ("    9.5 Consideraciones éticas", "19", False),
        ("18. Bibliografía (Normas APA 7.ª edición)", "20", True),
    ]
    
    table = doc.add_table(rows=len(items)+1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="102C57"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="102C57"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="dotted" w:sz="4" w:space="0" w:color="D0D7DE"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)
    
    row_hdr = table.rows[0]
    set_cell_background(row_hdr.cells[0], "F0F4F8")
    set_cell_background(row_hdr.cells[1], "F0F4F8")
    set_cell_margins(row_hdr.cells[0], top=60, bottom=60, left=100, right=100)
    set_cell_margins(row_hdr.cells[1], top=60, bottom=60, left=100, right=100)
    row_hdr.cells[0].width = Inches(5.8)
    row_hdr.cells[1].width = Inches(0.8)
    
    p_h0 = row_hdr.cells[0].paragraphs[0]
    p_h0.paragraph_format.space_before = Pt(0)
    p_h0.paragraph_format.space_after = Pt(0)
    r_h0 = p_h0.add_run("Contenido / Estructura del Anteproyecto")
    r_h0.bold = True
    r_h0.font.name = "Arial"
    r_h0.font.size = Pt(10)
    r_h0.font.color.rgb = RGBColor(16, 44, 87)
    
    p_h1 = row_hdr.cells[1].paragraphs[0]
    p_h1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_h1.paragraph_format.space_before = Pt(0)
    p_h1.paragraph_format.space_after = Pt(0)
    r_h1 = p_h1.add_run("Pág.")
    r_h1.bold = True
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(10)
    r_h1.font.color.rgb = RGBColor(16, 44, 87)
    
    for idx, (title, page, is_bold) in enumerate(items, start=1):
        row = table.rows[idx]
        cell_t = row.cells[0]
        cell_p = row.cells[1]
        cell_t.width = Inches(5.8)
        cell_p.width = Inches(0.8)
        set_cell_margins(cell_t, top=35, bottom=35, left=100, right=100)
        set_cell_margins(cell_p, top=35, bottom=35, left=100, right=100)
        
        pt = cell_t.paragraphs[0]
        pt.paragraph_format.space_before = Pt(0)
        pt.paragraph_format.space_after = Pt(0)
        pt.paragraph_format.line_spacing = 1.05
        rt = pt.add_run(title)
        rt.bold = is_bold
        rt.font.name = "Arial"
        rt.font.size = Pt(9.5 if not is_bold else 10)
        rt.font.color.rgb = RGBColor(16, 44, 87) if is_bold else RGBColor(40, 40, 40)
        
        pp = cell_p.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pp.paragraph_format.space_before = Pt(0)
        pp.paragraph_format.space_after = Pt(0)
        pp.paragraph_format.line_spacing = 1.05
        rp = pp.add_run(page)
        rp.bold = is_bold
        rp.font.name = "Arial"
        rp.font.size = Pt(9.5 if not is_bold else 10)
        rp.font.color.rgb = RGBColor(16, 44, 87) if is_bold else RGBColor(60, 60, 60)

def add_matrix_table(doc):
    headers = ["Categoría Central", "Subcategorías Temáticas", "Definición Conceptual y Operativa", "Evidencia y Fuentes Documentales"]
    rows_data = [
        (
            "1. Impacto Psicológico y Salud Mental",
            "• Depresión y sintomatología ansiosa\n• Estrés crónico e hipervigilancia\n• Trastorno de estrés postraumático (TEPT)\n• Sentimientos de culpa y desvalorización",
            "Afectaciones clínicas, subjetivas y emocionales derivadas de la exposición sostenida a situaciones de riesgo vital, violencia estructural, explotación y juicio social cotidiano en la zona de tolerancia.",
            "Escobar & Ramírez (2021); Nieto (2025); Figueroa-Ospina et al. (2023); Gutiérrez (2013); Ceballos-Ospino et al. (2013); Guillén (2022)."
        ),
        (
            "2. Estigma Social e Institucional",
            "• Estigma público y moralizante\n• Autoestigma e identidad deteriorada\n• Discriminación en la red hospitalaria\n• Manejo del secreto y la ocultación",
            "Proceso relacional de desacreditación social fundado en la etiqueta moral de 'mujer de la calle', que restringe el acceso equitativo a servicios públicos, legitima el rechazo familiar y fractura el autoconcepto.",
            "Goffman (1963/2009); Clemente Villar (2020); Mingorance (2015); Sierra Ospina et al. (2025); Álvarez (2021a); Sau (2000)."
        ),
        (
            "3. Condiciones de Vida y Precariedad",
            "• Informalidad laboral e inestabilidad\n• Violencia física, sexual y simbólica\n• Territorio urbano Comuna 10 (La Candelaria)\n• Dependencia del ingreso diario de subsistencia",
            "Factores materiales y estructurales del entorno urbano de Medellín que configuran escenarios de alta vulnerabilidad psicosocial y económica, condicionando el ejercicio forzado de la actividad.",
            "Caracol Televisión (2022); Cueto (2024); Godoy (2024); Montoya & Morales (2015); Cedeño et al. (2017); Bermúdez Pabón et al. (2007)."
        ),
        (
            "4. Estrategias de Afrontamiento y Resiliencia",
            "• Redes de solidaridad y cuidado entre pares\n• Resignificación del rol materno/familiar\n• Afrontamiento activo vs. paliativo/evitativo\n• Espacios de escucha y empoderamiento",
            "Mecanismos psicológicos, relacionales y comunitarios orientados a regular el sufrimiento, amortiguar las cargas emocionales y construir proyectos de vida dignos en medio de la adversidad.",
            "Lazarus & Folkman (1984); Mayordomo et al. (2016); García & Rodríguez (2019); González et al. (2018); Rodríguez (2014); Rueda (2020)."
        ),
        (
            "5. Políticas Públicas y Respuesta Gubernamental (2020–2024)",
            "• Cobertura del programa 'Por Mis Derechos'\n• Política Pública de Salud Mental (Acuerdo 043/16)\n• Sesgo biomédico frente al cuidado psicosocial\n• Rutas de atención y mandatos jurisprudenciales",
            "Conjunto de directrices, normativas y acciones estatales implementadas por el Distrito de Medellín y el Gobierno Nacional para garantizar la atención en salud mental y derechos fundamentales.",
            "Alcaldía de Medellín (Gómez, 2020); Sentencias T-629/10 y T-594/16; Ley 1616/13; Médicos Sin Fronteras (2023); MinSalud (2025); MinIgualdad (2024)."
        )
    ]
    
    table = doc.add_table(rows=len(rows_data)+1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="102C57", sz="4")
    
    col_widths = [Inches(1.6), Inches(1.7), Inches(2.0), Inches(1.5)]
    
    hdr = table.rows[0]
    make_row_cant_split(hdr)
    for c_idx, title in enumerate(headers):
        cell = hdr.cells[c_idx]
        cell.width = col_widths[c_idx]
        set_cell_background(cell, "102C57")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, (cat, subcat, defin, fuentes) in enumerate(rows_data, start=1):
        row = table.rows[r_idx]
        make_row_cant_split(row)
        bg = "FFFFFF" if r_idx % 2 != 0 else "F7FAFC"
        data_cells = [cat, subcat, defin, fuentes]
        for c_idx, val in enumerate(data_cells):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            if bg != "FFFFFF":
                set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(16, 44, 87)
            else:
                r.font.color.rgb = RGBColor(30, 30, 30)

print("make_anteproyecto_docx.py ready.")
