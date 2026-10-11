# -*- coding: utf-8 -*-
import docx

doc = docx.Document(r"C:\Users\lenovo\Downloads\Plantilla Tipo Tarea_CIENCIAS SOCIALES Y AMBIENTE -ACTIVIDAD-1.docx")

print(f"Paragraphs count: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    print(f"P{i}: '{p.text}'")

print(f"\nTables count: {len(doc.tables)}")
for t_i, t in enumerate(doc.tables):
    print(f"\n=== Table {t_i}: rows={len(t.rows)}, cols={len(t.columns)} ===")
    for r_i, r in enumerate(t.rows):
        for c_i, c in enumerate(r.cells):
            print(f"T{t_i} R{r_i} C{c_i}: (len={len(c.text)})")
            lines = [l.strip() for l in c.text.split('\n') if l.strip()]
            for l in lines[:10]:
                print(f"    {l[:100]}")
            if len(lines) > 10:
                print(f"    ... [{len(lines)-10} more lines] ...")
            # check nested tables in this cell
            for n_i, nt in enumerate(c.tables):
                print(f"    [NESTED TABLE {n_i}: rows={len(nt.rows)}, cols={len(nt.columns)}]")
                for nr_i, nr in enumerate(nt.rows):
                    print(f"      NR{nr_i}: {[nc.text.strip().replace(chr(10), ' ') for nc in nr.cells]}")
