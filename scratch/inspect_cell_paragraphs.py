# -*- coding: utf-8 -*-
import docx

doc = docx.Document(r"C:\Users\lenovo\Downloads\Plantilla Tipo Tarea_CIENCIAS SOCIALES Y AMBIENTE -ACTIVIDAD-1.docx")

t1 = doc.tables[1]
cell = t1.rows[2].cells[0]
print("--- ALL PARAGRAPHS IN TABLE 1 ROW 2 CELL 0 ---")
for i, p in enumerate(cell.paragraphs):
    print(f"p{i}: {repr(p.text)}")
