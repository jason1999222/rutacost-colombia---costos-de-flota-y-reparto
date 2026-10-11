import docx

doc = docx.Document(r"C:\Users\lenovo\Downloads\Plantilla Tipo Tarea_CIENCIAS SOCIALES Y AMBIENTE -ACTIVIDAD-1.docx")
cell = doc.tables[1].rows[2].cells[0]

print("=== MUESTREO DE ENCABEZADOS Y FORMATO ===")
for i, p in enumerate(cell.paragraphs):
    txt = p.text.strip()
    if any(h in txt for h in ['Fase 1:', '1. ¿Cuál es el problema', '2. ¿Cuáles son sus', 'Fase 2:', 'Fase 3:', 'Fase 4:', 'Referencias Bibliográficas']):
        runs = p.runs
        bolds = [r.bold for r in runs]
        colors = [str(r.font.color.rgb) if r.font.color and r.font.color.rgb else "None" for r in runs]
        print(f"P{i:02d}: {repr(txt[:40])} | Bold={bolds} | Color={colors}")
