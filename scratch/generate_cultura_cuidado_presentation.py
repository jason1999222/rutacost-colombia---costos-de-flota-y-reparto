# -*- coding: utf-8 -*-
r"""
Generador de Presentación de Alta Calidad Estilo Canva:
'LA CULTURA DEL CUIDADO'
Presentado por: Bertha Inés Medina Hernández

Características:
- Formato 16:9 Panorámico (Widescreen 13.333" x 7.5").
- Paleta de diseño sofisticada y moderna (Slate Dark #0F172A, Emerald/Teal #0D9488, Warm Coral #E11D48, Amber Gold #D97706, Light #F8FAFC).
- Tarjetas con bordes limpios, etiquetas redondeadas (badges), jerarquía tipográfica premium.
- Contenido profundo, inspirador y riguroso sobre las 3 dimensiones del cuidado (autocuidado, cuidado del otro y ecocuidado).
- Exportable a PDF y 100% compatible con la importación directa en Canva (.pptx editable).
"""

import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Paleta de Colores
C_DARK_BG   = RGBColor(15, 23, 42)      # Slate 900
C_DARK_CARD = RGBColor(30, 41, 59)      # Slate 800
C_LIGHT_BG  = RGBColor(248, 250, 252)   # Slate 50
C_WHITE     = RGBColor(255, 255, 255)
C_TEXT_DARK = RGBColor(30, 41, 59)      # Slate 800
C_TEXT_MUTED= RGBColor(100, 116, 139)   # Slate 500
C_TEXT_LIGHT= RGBColor(241, 245, 249)   # Slate 100
C_TEAL      = RGBColor(13, 148, 136)    # Teal 600
C_TEAL_DARK = RGBColor(15, 118, 110)    # Teal 700
C_TEAL_LIGHT= RGBColor(204, 251, 241)   # Teal 100
C_CORAL     = RGBColor(225, 29, 72)     # Rose 600
C_AMBER     = RGBColor(217, 119, 6)     # Amber 600
C_BORDER    = RGBColor(226, 232, 240)   # Slate 200

FONT_MAIN = "Arial"

def set_slide_bg(slide, prs, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_badge(slide, text, left, top, bg_color=C_TEAL, text_color=C_WHITE, font_size=10):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.2), Inches(0.35))
    badge.fill.solid()
    badge.fill.fore_color.rgb = bg_color
    badge.line.fill.background()
    tf = badge.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT_MAIN
    r.font.size = Pt(font_size)
    r.font.bold = True
    r.font.color.rgb = text_color
    return badge

def add_header(slide, badge_text, title_text, subtitle_text=None, is_dark=False):
    left = Inches(0.8)
    top = Inches(0.6)
    badge_bg = C_TEAL if not is_dark else C_TEAL_DARK
    badge = add_badge(slide, badge_text, left, top, bg_color=badge_bg, text_color=C_WHITE)
    
    # Title
    t_box = slide.shapes.add_textbox(left, Inches(1.05), Inches(11.7), Inches(0.7))
    tf = t_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title_text
    r.font.name = FONT_MAIN
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = C_WHITE if is_dark else C_TEXT_DARK
    
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(left, Inches(1.75), Inches(11.7), Inches(0.45))
        stf = sub_box.text_frame
        stf.word_wrap = True
        stf.margin_left = stf.margin_top = stf.margin_right = stf.margin_bottom = 0
        sp = stf.paragraphs[0]
        sr = sp.add_run()
        sr.text = subtitle_text
        sr.font.name = FONT_MAIN
        sr.font.size = Pt(13)
        sr.font.color.rgb = C_TEXT_LIGHT if is_dark else C_TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=C_WHITE, border_color=C_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
    else:
        card.line.fill.background()
    return card

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ==========================================
    # SLIDE 1: PORTADA PRINCIPAL (DARK THEME)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, prs, C_DARK_BG)
    
    # Decorative accent banner
    accent_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(0.12), Inches(5.5))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = C_TEAL
    accent_bar.line.fill.background()

    # Pill badge
    add_badge(s1, "ÉTICA, BIENESTAR Y SOSTENIBILIDAD", Inches(1.2), Inches(1.2), bg_color=C_TEAL, font_size=10.5)

    # Big Title
    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(2.2))
    tf = t_box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "LA CULTURA\nDEL CUIDADO"
    r1.font.name = FONT_MAIN
    r1.font.size = Pt(44)
    r1.font.bold = True
    r1.font.color.rgb = C_WHITE

    p2 = tf.add_paragraph()
    p2.space_before = Pt(16)
    r2 = p2.add_run()
    r2.text = "Una propuesta para la preservación de la vida, las relaciones humanas y la corresponsabilidad colectiva"
    r2.font.name = FONT_MAIN
    r2.font.size = Pt(17)
    r2.font.color.rgb = RGBColor(203, 213, 225)

    # Author card at bottom
    author_card = add_card(s1, Inches(1.2), Inches(4.7), Inches(10.5), Inches(1.8), bg_color=C_DARK_CARD, border_color=C_TEAL)
    a_box = s1.shapes.add_textbox(Inches(1.5), Inches(4.85), Inches(9.9), Inches(1.5))
    atf = a_box.text_frame
    atf.word_wrap = True
    ap1 = atf.paragraphs[0]
    ar1 = ap1.add_run()
    ar1.text = "PRESENTADO POR:"
    ar1.font.size = Pt(11)
    ar1.font.bold = True
    ar1.font.color.rgb = C_TEAL

    ap2 = atf.add_paragraph()
    ap2.space_before = Pt(4)
    ar2 = ap2.add_run()
    ar2.text = "Bertha Inés Medina Hernández"
    ar2.font.size = Pt(22)
    ar2.font.bold = True
    ar2.font.color.rgb = C_WHITE

    ap3 = atf.add_paragraph()
    ap3.space_before = Pt(4)
    ar3 = ap3.add_run()
    ar3.text = "Cátedra de Ciencias Sociales, Ética y Desarrollo Humano • 2026"
    ar3.font.size = Pt(12)
    ar3.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # SLIDE 2: ¿QUÉ ES LA CULTURA DEL CUIDADO?
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2, prs, C_LIGHT_BG)
    add_header(s2, "CONCEPTOS CLAVE", "¿Qué es la Cultura del Cuidado?", "Fundamentos éticos para comprender el cuidado como soporte de la vida")

    col_w = Inches(3.64)
    col_h = Inches(4.7)
    cards_data_s2 = [
        ("01", "Modo Esencial de Ser", C_TEAL,
         "Según el filósofo Leonardo Boff, el cuidado no es un acto aislado ni una técnica accesoria; es la estructura fundamental que nos hace humanos. Sin cuidado, el ser humano se deshumaniza y perece desde su nacimiento.",
         "• Dimensión ontológica del ser.\n• Más que una tarea: una actitud vital.\n• Conexión afectiva con el entorno."),
        
        ("02", "Superar la Indiferencia", C_CORAL,
         "Vivimos en sociedades apresuradas que normalizan el descarte y la soledad. La cultura del cuidado rompe la ceguera cotidiana y nos invita a prestar atención consciente a la vulnerabilidad propia y de quienes nos rodean.",
         "• Atención activa y sensible.\n• Desarticulación del egoísmo.\n• Reconocimiento de la fragilidad."),
        
        ("03", "Corresponsabilidad Viva", C_AMBER,
         "El cuidado no puede recaer exclusivamente en una persona, en el médico o en el personal de limpieza. Constituye un pacto social donde cada miembro asume su cuota en la preservación de la armonía comunitaria.",
         "• Deber cívico compartido.\n• Acción recíproca y solidaria.\n• Sostenibilidad del tejido social.")
    ]

    for idx, (num, title, color, desc, points) in enumerate(cards_data_s2):
        l = Inches(0.8 + idx * (3.64 + 0.39))
        add_card(s2, l, Inches(2.2), col_w, col_h, bg_color=C_WHITE, border_color=C_BORDER)
        
        # Color bar top
        bar = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, Inches(2.2), col_w, Inches(0.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tb = s2.shapes.add_textbox(l + Inches(0.3), Inches(2.45), col_w - Inches(0.6), col_h - Inches(0.55))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_num = tf.paragraphs[0]
        r_num = p_num.add_run()
        r_num.text = num
        r_num.font.size = Pt(28)
        r_num.font.bold = True
        r_num.font.color.rgb = color

        p_t = tf.add_paragraph()
        p_t.space_before = Pt(4)
        r_t = p_t.add_run()
        r_t.text = title
        r_t.font.size = Pt(16)
        r_t.font.bold = True
        r_t.font.color.rgb = C_TEXT_DARK

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(8)
        r_d = p_d.add_run()
        r_d.text = desc
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = C_TEXT_DARK

        p_pts = tf.add_paragraph()
        p_pts.space_before = Pt(12)
        r_pts = p_pts.add_run()
        r_pts.text = points
        r_pts.font.size = Pt(10.5)
        r_pts.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 3: EL TRIÁNGULO DEL CUIDADO (3 DIMENSIONES)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3, prs, C_LIGHT_BG)
    add_header(s3, "MODELO SISTÉMICO", "Las Tres Dimensiones del Cuidado", "El cuidado integral opera en tres planos inseparables y mutuamente dependientes")

    dims = [
        ("AUTOCUIDADO", "Cuidado de Sí Mismo", C_TEAL,
         "La raíz indispensable. Es la capacidad de reconocer nuestras propias necesidades físicas, mentales y emocionales. No es narcisismo ni egoísmo: es la base que permite estar en condiciones de brindar apoyo al prójimo.",
         "Claves prácticas:\n• Sueño reparador y alimentación consciente.\n• Poner límites claros y aprender a decir 'no'.\n• Gestión del estrés y pausas restaurativas."),

        ("CUIDADO DEL OTRO", "Relaciones e Interdependencia", C_CORAL,
         "La dimensión vincular. Se apoya en la empatía, la compasión y el reconocimiento del prójimo como un legítimo otro. Se fundamenta en la Ética del Cuidado de Carol Gilligan, situando los afectos y la dignidad en el centro.",
         "Claves prácticas:\n• Escucha atenta sin juzgar apresuradamente.\n• Buen trato verbal y comunicación asertiva.\n• Apoyo mutuo en momentos de dificultad."),

        ("ECOCUIDADO", "Cuidado del Entorno y la Naturaleza", C_AMBER,
         "La dimensión ecológica y comunitaria. Reconoce que habitamos una 'Casa Común' finita y vulnerable. Implica asumir hábitos que protejan el hábitat compartido y garanticen un futuro digno para las próximas generaciones.",
         "Claves prácticas:\n• Consumo responsable y reducción de residuos.\n• Protección de las zonas verdes y el agua.\n• Participación en iniciativas socioambientales.")
    ]

    for idx, (badge, title, color, desc, points) in enumerate(dims):
        l = Inches(0.8 + idx * (3.64 + 0.39))
        add_card(s3, l, Inches(2.2), col_w, col_h, bg_color=C_WHITE, border_color=C_BORDER)

        add_badge(s3, badge, l + Inches(0.3), Inches(2.45), bg_color=color, font_size=9.5)

        tb = s3.shapes.add_textbox(l + Inches(0.3), Inches(2.95), col_w - Inches(0.6), col_h - Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title
        r_t.font.size = Pt(15)
        r_t.font.bold = True
        r_t.font.color.rgb = C_TEXT_DARK

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(8)
        r_d = p_d.add_run()
        r_d.text = desc
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = C_TEXT_DARK

        p_pts = tf.add_paragraph()
        p_pts.space_before = Pt(12)
        r_pts = p_pts.add_run()
        r_pts.text = points
        r_pts.font.size = Pt(10)
        r_pts.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 4: DIMENSIÓN 1 - AUTOCUIDADO
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4, prs, C_LIGHT_BG)
    add_header(s4, "DIMENSIÓN 1", "El Autocuidado: Cuidarse para Trascender", "«Nadie puede ofrecer bienestar si su propia vida se encuentra agotada o desatendida»")

    # Banner quote
    q_card = add_card(s4, Inches(0.8), Inches(2.2), Inches(11.7), Inches(0.8), bg_color=C_TEAL_LIGHT, border_color=C_TEAL)
    q_tb = s4.shapes.add_textbox(Inches(1.1), Inches(2.25), Inches(11.1), Inches(0.7))
    q_tf = q_tb.text_frame
    q_tf.word_wrap = True
    qp = q_tf.paragraphs[0]
    qr = qp.add_run()
    qr.text = "Principio Rector: El autocuidado no es un lujo ni un acto de vanidad; es una responsabilidad ética primaria para preservar la salud integral."
    qr.font.bold = True
    qr.font.size = Pt(12)
    qr.font.color.rgb = C_TEAL_DARK

    # 4 Cards in 2x2 grid
    grid_w = Inches(5.65)
    grid_h = Inches(1.85)
    auto_items = [
        ("1. Salud Física y Vitalidad", "Garantizar descanso suficiente (7-8 horas), alimentación equilibrada y pausas activas durante jornadas prolongadas. El cuerpo avisa antes de colapsar.", Inches(0.8), Inches(3.2)),
        ("2. Gestión Emocional Consciente", "Validar lo que sentimos: el cansancio, la tristeza o la frustración no son debilidades. Aprender a procesar las emociones sin reprimirlas ni desbordarse.", Inches(6.88), Inches(3.2)),
        ("3. Establecimiento de Límites", "El poder restaurativo de decir 'NO'. Proteger el tiempo personal frente a la hiperconexión digital y evitar sobrecargarse de compromisos ajenos.", Inches(0.8), Inches(5.25)),
        ("4. Autoestima y Autocompasión", "Tratarse a uno mismo con la misma bondad y paciencia que solemos brindar a un buen amigo. Erradicar la autocrítica destructiva y el perfeccionismo tóxico.", Inches(6.88), Inches(5.25))
    ]

    for title, desc, left, top in auto_items:
        add_card(s4, left, top, grid_w, grid_h, bg_color=C_WHITE, border_color=C_BORDER)
        tb = s4.shapes.add_textbox(left + Inches(0.3), top + Inches(0.2), grid_w - Inches(0.6), grid_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title
        r_t.font.size = Pt(14)
        r_t.font.bold = True
        r_t.font.color.rgb = C_TEAL

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(4)
        r_d = p_d.add_run()
        r_d.text = desc
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = C_TEXT_DARK

    # ==========================================
    # SLIDE 5: DIMENSIÓN 2 - CUIDADO DEL OTRO
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5, prs, C_LIGHT_BG)
    add_header(s5, "DIMENSIÓN 2", "El Cuidado del Otro: Ética del Buen Trato", "Construir comunidades compasivas donde la dignidad humana sea innegociable")

    cards_otro = [
        ("Escucha Activa y Presencia Genuina", C_CORAL,
         "En un mundo hiperconectado pero incomunicado, escuchar sin interrumpir es un acto de cuidado revolucionario. Estar presente con el cuerpo y la mente para validar la vivencia del otro."),
        
        ("Comunicación Asertiva y Cuidado de la Palabra", C_TEAL,
         "Las palabras tienen poder: pueden construir puentes o levantar muros. La cultura del cuidado exige desterrar el maltrato sutil, el sarcasmo hiriente y la descalificación en nuestras conversaciones diarias."),
        
        ("Reconocimiento de la Economía del Cuidado", C_AMBER,
         "Históricamente, las labores de crianza, alimentación y apoyo emocional han recaído sobre las mujeres de forma no remunerada. Cuidar al otro exige democratizar y valorar justamente estas tareas."),
        
        ("Solidaridad en Momentos de Vulnerabilidad", C_CORAL,
         "Acompañar en el duelo, la enfermedad o la crisis económica. La red comunitaria de apoyo mutuo es el verdadero escudo frente a la incertidumbre y el sufrimiento.")
    ]

    for idx, (title, color, text) in enumerate(cards_otro):
        l = Inches(0.8 + idx * (2.7 + 0.3))
        add_card(s5, l, Inches(2.2), Inches(2.7), Inches(4.7), bg_color=C_WHITE, border_color=C_BORDER)

        bar = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, Inches(2.2), Inches(2.7), Inches(0.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tb = s5.shapes.add_textbox(l + Inches(0.25), Inches(2.5), Inches(2.2), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title
        r_t.font.size = Pt(14)
        r_t.font.bold = True
        r_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(10)
        r_d = p_d.add_run()
        r_d.text = text
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = C_TEXT_DARK

    # ==========================================
    # SLIDE 6: DIMENSIÓN 3 - ECOCUIDADO
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6, prs, C_LIGHT_BG)
    add_header(s6, "DIMENSIÓN 3", "El Cuidado del Entorno: Nuestra Casa Común", "Armonía entre la actividad humana y los límites biofísicos de la naturaleza")

    eco_cards = [
        ("Consumo Consciente y Sobriedad",
         "Cuestionar la necesidad real de lo que adquirimos. Reducir plásticos de un solo uso, combatir el desperdicio de comida y preferir productos locales y duraderos."),
        
        ("Corresponsabilidad con el Territorio",
         "El espacio público no es de 'nadie': es de todos. Separar residuos en la fuente, cuidar las fuentes de agua y valorar la labor fundamental de los recicladores de oficio."),
        
        ("Justicia Intergeneracional",
         "La ética del cuidado mira hacia adelante: ¿qué mundo le heredaremos a las futuras generaciones? Actuar hoy para frenar la pérdida de biodiversidad y el cambio climático.")
    ]

    for idx, (title, desc) in enumerate(eco_cards):
        l = Inches(0.8 + idx * (3.64 + 0.39))
        add_card(s6, l, Inches(2.2), col_w, Inches(4.7), bg_color=C_WHITE, border_color=C_BORDER)

        # Header box in card
        hbar = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l + Inches(0.3), Inches(2.5), col_w - Inches(0.6), Inches(0.8))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = C_TEAL_LIGHT
        hbar.line.fill.background()
        htf = hbar.text_frame
        htf.vertical_anchor = MSO_ANCHOR.MIDDLE
        hp = htf.paragraphs[0]
        hp.alignment = PP_ALIGN.CENTER
        hr = hp.add_run()
        hr.text = title
        hr.font.size = Pt(13)
        hr.font.bold = True
        hr.font.color.rgb = C_TEAL_DARK

        tb = s6.shapes.add_textbox(l + Inches(0.3), Inches(3.5), col_w - Inches(0.6), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = desc
        r.font.size = Pt(11.5)
        r.font.color.rgb = C_TEXT_DARK

    # ==========================================
    # SLIDE 7: EN EL ÁMBITO LABORAL E INSTITUCIONAL
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7, prs, C_LIGHT_BG)
    add_header(s7, "ORGANIZACIONES SALUDABLES", "La Cultura del Cuidado en el Trabajo", "De la prevención obligatoria a un compromiso colectivo por la vida")

    work_items = [
        ("Seguridad y Salud en el Trabajo (SST)",
         "La prevención de riesgos no es un conjunto de formatos archivados: es la certeza de que cada persona regrese sana y salva a su hogar. Exige ergonomía, orden y protección activa."),
        
        ("Ambientes Libres de Acoso y Hostilidad",
         "La salud mental laboral se destruye ante el maltrato o la sobrecarga desmedida. Las organizaciones del cuidado promueven la seguridad psicológica para opinar y equivocarse."),
        
        ("Liderazgo Cercano y Empático",
         "Líderes que preguntan '¿cómo estás hoy?' antes de pedir resultados numéricos. La empatía fomenta mayor compromiso, productividad y lealtad que la presión autoritaria."),
        
        ("Cuidado Mutuo entre Colegas",
         "Si observamos a un compañero sobrecargado, enfermo o cometiendo un acto inseguro, intervenimos con respeto. Cuidar al colega es proteger la seguridad de todo el equipo.")
    ]

    for idx, (title, desc) in enumerate(work_items):
        r_idx = idx // 2
        c_idx = idx % 2
        l = Inches(0.8 + c_idx * (5.65 + 0.4))
        t = Inches(2.2 + r_idx * (2.25 + 0.25))

        add_card(s7, l, t, Inches(5.65), Inches(2.25), bg_color=C_WHITE, border_color=C_BORDER)
        
        # Left color strip
        strip = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, Inches(0.12), Inches(2.25))
        strip.fill.solid()
        strip.fill.fore_color.rgb = C_TEAL if c_idx == 0 else C_AMBER
        strip.line.fill.background()

        tb = s7.shapes.add_textbox(l + Inches(0.3), t + Inches(0.2), Inches(5.1), Inches(1.85))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title
        r_t.font.size = Pt(14)
        r_t.font.bold = True
        r_t.font.color.rgb = C_TEXT_DARK

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(6)
        r_d = p_d.add_run()
        r_d.text = desc
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = C_TEXT_DARK

    # ==========================================
    # SLIDE 8: BARRERAS Y ENEMIGOS DEL CUIDADO
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8, prs, C_LIGHT_BG)
    add_header(s8, "DIAGNÓSTICO CRÍTICO", "¿Qué Obstaculiza la Cultura del Cuidado?", "Identificar las trampas cotidianas que fracturan nuestra capacidad de cuidar")

    barreras = [
        ("1. Individualismo y Apatía", C_CORAL,
         "La mentalidad del 'sálvese quien pueda'. Desentenderse del sufrimiento o la sobrecarga ajena asumiendo que 'no es mi problema'."),
        
        ("2. Culto a la Prisa y al Agotamiento", C_AMBER,
         "Normalizar el estrés extremo y la falta de sueño como si fueran medallas de éxito profesional. La prisa destruye la empatía."),
        
        ("3. Desentendimiento Moral", C_CORAL,
         "Pensar 'para eso le pagan a otros'. Justificar el desorden o la falta de respeto delegando la responsabilidad en terceros."),
        
        ("4. Precarización de Quienes Cuidan", C_AMBER,
         "Invisibilizar a las personas que limpian, curan, cocinan y cuidan niños o ancianos, negándoles salario digno y reconocimiento social.")
    ]

    for idx, (title, color, text) in enumerate(barreras):
        l = Inches(0.8 + idx * (2.7 + 0.3))
        add_card(s8, l, Inches(2.2), Inches(2.7), Inches(4.7), bg_color=C_WHITE, border_color=C_BORDER)

        bar = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, Inches(2.2), Inches(2.7), Inches(0.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        tb = s8.shapes.add_textbox(l + Inches(0.25), Inches(2.5), Inches(2.2), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title
        r_t.font.size = Pt(14)
        r_t.font.bold = True
        r_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(10)
        r_d = p_d.add_run()
        r_d.text = text
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = C_TEXT_DARK

    # ==========================================
    # SLIDE 9: ESTRATEGIAS PRÁCTICAS COTIDIANAS
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9, prs, C_LIGHT_BG)
    add_header(s9, "HERRAMIENTAS DE ACCIÓN", "Estrategias para Vivir la Cultura del Cuidado", "Pequeñas acciones cotidianas que generan grandes transformaciones colectivas")

    cols_est = [
        ("Nivel Personal", C_TEAL,
         "• Pausas conscientes de respiración al día.\n• Establecer horarios de desconexión nocturna.\n• Hidratarse y consumir alimentos nutritivos.\n• Expresar emociones sin juzgarse duramente."),
        
        ("Nivel Relacional", C_CORAL,
         "• Saludar con calidez y mirar a los ojos.\n• Preguntar genuinamente '¿cómo te sientes?'.\n• Practicar el agradecimiento explícito.\n• Brindar ayuda sin esperar retribución."),
        
        ("Nivel Comunitario", C_AMBER,
         "• Respetar normas de convivencia y silencio.\n• Cuidar las instalaciones y bienes comunes.\n• Separar residuos y apoyar al reciclador.\n• Participar en redes y comités de apoyo social.")
    ]

    for idx, (title, color, text) in enumerate(cols_est):
        l = Inches(0.8 + idx * (3.64 + 0.39))
        add_card(s9, l, Inches(2.2), col_w, Inches(4.7), bg_color=C_WHITE, border_color=C_BORDER)

        hbar = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l + Inches(0.3), Inches(2.5), col_w - Inches(0.6), Inches(0.7))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = color
        hbar.line.fill.background()
        htf = hbar.text_frame
        htf.vertical_anchor = MSO_ANCHOR.MIDDLE
        hp = htf.paragraphs[0]
        hp.alignment = PP_ALIGN.CENTER
        hr = hp.add_run()
        hr.text = title
        hr.font.size = Pt(14)
        hr.font.bold = True
        hr.font.color.rgb = C_WHITE

        tb = s9.shapes.add_textbox(l + Inches(0.4), Inches(3.4), col_w - Inches(0.8), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = text
        r.font.size = Pt(12)
        r.font.color.rgb = C_TEXT_DARK

    # ==========================================
    # SLIDE 10: DECÁLOGO DE LA CULTURA DEL CUIDADO
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10, prs, C_LIGHT_BG)
    add_header(s10, "PRINCIPIOS GUÍA", "Decálogo de la Cultura del Cuidado", "Diez compromisos esenciales para cultivar en el día a día")

    decalogo_col1 = [
        ("1. La vida es sagrada:", "Proteger la integridad física, mental y espiritual de todo ser vivo."),
        ("2. El autocuidado es la base:", "Cuidarse a sí mismo es el primer requisito para cuidar a otros."),
        ("3. La palabra edifica:", "Usar un lenguaje respetuoso, constructivo y libre de agresiones."),
        ("4. Escuchar es abrazar:", "Prestar atención plena sin juzgar ni apresurar conclusiones."),
        ("5. La vulnerabilidad une:", "Aceptar que todos necesitamos ayuda en algún momento de la vida.")
    ]
    decalogo_col2 = [
        ("6. Corresponsabilidad activa:", "Asumir las consecuencias de nuestras acciones y desidias."),
        ("7. Gratitud y reconocimiento:", "Valorar y dignificar a quienes realizan tareas de cuidado."),
        ("8. Habitar la Tierra con respeto:", "Consumir con sobriedad y no agotar los recursos naturales."),
        ("9. Paz en el conflicto:", "Resolver diferencias mediante el diálogo sincero y no violento."),
        ("10. Cuidar es trascender:", "Dejar una huella de bienestar y amor por donde pasemos.")
    ]

    card_d1 = add_card(s10, Inches(0.8), Inches(2.2), Inches(5.65), Inches(4.7), bg_color=C_WHITE, border_color=C_BORDER)
    card_d2 = add_card(s10, Inches(6.88), Inches(2.2), Inches(5.65), Inches(4.7), bg_color=C_WHITE, border_color=C_BORDER)

    def populate_decalogo(box, items):
        tf = box.text_frame
        tf.word_wrap = True
        for i, (title, desc) in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(8) if i > 0 else Pt(0)
            rt = p.add_run()
            rt.text = title + " "
            rt.font.bold = True
            rt.font.size = Pt(11)
            rt.font.color.rgb = C_TEAL
            
            rd = p.add_run()
            rd.text = desc
            rd.font.size = Pt(10.5)
            rd.font.color.rgb = C_TEXT_DARK

    tb1 = s10.shapes.add_textbox(Inches(1.1), Inches(2.35), Inches(5.05), Inches(4.4))
    populate_decalogo(tb1, decalogo_col1)

    tb2 = s10.shapes.add_textbox(Inches(7.18), Inches(2.35), Inches(5.05), Inches(4.4))
    populate_decalogo(tb2, decalogo_col2)

    # ==========================================
    # SLIDE 11: CONCLUSIÓN Y SÍNTESIS (DARK THEME)
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11, prs, C_DARK_BG)
    add_header(s11, "REFLEXIÓN FINAL", "El Cuidado como Fuerza Transformadora", "Hacia una sociedad que pone la vida y el corazón en el centro", is_dark=True)

    c_box = s11.shapes.add_textbox(Inches(1.5), Inches(2.6), Inches(10.3), Inches(2.6))
    ctf = c_box.text_frame
    ctf.word_wrap = True
    
    cp1 = ctf.paragraphs[0]
    cr1 = cp1.add_run()
    cr1.text = "«El cuidado no es una opción secundaria ni un gesto caritativo ocasional: constituye la condición irrenunciable para que la sociedad no se desmorone. Cuidar de nosotros mismos, abrazar la fragilidad del otro y sanar nuestra relación con la Tierra es el mayor acto de madurez ética al que podemos aspirar.»"
    cr1.font.name = FONT_MAIN
    cr1.font.size = Pt(20)
    cr1.font.italic = True
    cr1.font.color.rgb = C_WHITE

    # Author card
    aut_c = add_card(s11, Inches(1.5), Inches(5.2), Inches(10.3), Inches(1.2), bg_color=C_DARK_CARD, border_color=C_TEAL)
    at_box = s11.shapes.add_textbox(Inches(1.8), Inches(5.35), Inches(9.7), Inches(0.9))
    attf = at_box.text_frame
    attf.word_wrap = True
    ap = attf.paragraphs[0]
    ar = ap.add_run()
    ar.text = "Bertha Inés Medina Hernández"
    ar.font.bold = True
    ar.font.size = Pt(18)
    ar.font.color.rgb = C_WHITE

    ap2 = attf.add_paragraph()
    ap2.space_before = Pt(3)
    ar2 = ap2.add_run()
    ar2.text = "La Cultura del Cuidado • Un compromiso ético para vivir con dignidad y esperanza"
    ar2.font.size = Pt(12)
    ar2.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # SLIDE 12: CIERRE Y PREGUNTAS (DARK THEME)
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12, prs, C_DARK_BG)

    cierre_box = s12.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.3), Inches(4.0))
    ztf = cierre_box.text_frame
    ztf.word_wrap = True
    
    zp1 = ztf.paragraphs[0]
    zp1.alignment = PP_ALIGN.CENTER
    zr1 = zp1.add_run()
    zr1.text = "¡Muchas Gracias!"
    zr1.font.name = FONT_MAIN
    zr1.font.size = Pt(44)
    zr1.font.bold = True
    zr1.font.color.rgb = C_WHITE

    zp2 = ztf.add_paragraph()
    zp2.space_before = Pt(14)
    zp2.alignment = PP_ALIGN.CENTER
    zr2 = zp2.add_run()
    zr2.text = "«Quien cuida, transforma su entorno y deja una huella imborrable en la vida de los demás»"
    zr2.font.name = FONT_MAIN
    zr2.font.size = Pt(17)
    zr2.font.italic = True
    zr2.font.color.rgb = RGBColor(203, 213, 225)

    zp3 = ztf.add_paragraph()
    zp3.space_before = Pt(28)
    zp3.alignment = PP_ALIGN.CENTER
    zr3 = zp3.add_run()
    zr3.text = "Bertha Inés Medina Hernández"
    zr3.font.name = FONT_MAIN
    zr3.font.size = Pt(20)
    zr3.font.bold = True
    zr3.font.color.rgb = C_TEAL

    zp4 = ztf.add_paragraph()
    zp4.space_before = Pt(8)
    zp4.alignment = PP_ALIGN.CENTER
    zr4 = zp4.add_run()
    zr4.text = "Espacio abierto para reflexiones, preguntas y aportes"
    zr4.font.name = FONT_MAIN
    zr4.font.size = Pt(13)
    zr4.font.color.rgb = RGBColor(148, 163, 184)

    # Guardar en las rutas requeridas
    p_desktop = r"C:\Users\lenovo\Desktop\trabajos andres\Presentacion_Cultura_del_Cuidado_Bertha_Medina.pptx"
    p_downloads = r"C:\Users\lenovo\Downloads\Presentacion_Cultura_del_Cuidado_Bertha_Medina.pptx"

    print("Guardando presentación en Escritorio...")
    prs.save(p_desktop)
    print("Guardando presentación en Descargas...")
    prs.save(p_downloads)
    print("Presentación PPTX generada exitosamente con 12 diapositivas premium.")

if __name__ == "__main__":
    create_presentation()
